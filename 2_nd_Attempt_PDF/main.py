import tkinter as tk
import pandas as pd
import shutil
from tkinter import filedialog, messagebox, ttk
import ttkbootstrap as tb
import os
import logging
import darkdetect
import sys
from datetime import datetime
from docx import Document
from file_reader import read_excel_csv
from data_mapper import scan_template_placeholders, prepare_row_data, replace_all_placeholders
from docx2pdf import convert

# Configure logging
logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")


class DocumentFillerApp:
    def __init__(self, root):
        self.root = root
        self.load_default_templates()  # Load templates first
        self.setup_ui()
        self.setup_menu()

        # Initialize variables
        self.input_file = None
        self.output_folder = None
        self.current_data = None

        self.setup_template_access()
    def load_default_templates(self):
        """Load default templates from the templates folder"""
        try:
            # Get the directory where the executable or script is located
            if getattr(sys, 'frozen', False):
                # Running as compiled executable
                application_path = os.path.dirname(sys.executable)
            else:
                # Running as script
                application_path = os.path.dirname(os.path.abspath(__file__))

            templates_dir = os.path.join(application_path, "templates")

            self.eligible_template = os.path.join(templates_dir, "eligible_template.docx")
            self.ineligible_template = os.path.join(templates_dir, "ineligible_template.docx")

            if not os.path.exists(self.eligible_template):
                raise FileNotFoundError(f"Eligible template not found at {self.eligible_template}")
            if not os.path.exists(self.ineligible_template):
                raise FileNotFoundError(f"Ineligible template not found at {self.ineligible_template}")

            logging.info("Default templates loaded successfully")

        except Exception as e:
            logging.error(f"Failed to load default templates: {str(e)}")
            messagebox.showerror("Error", f"Failed to load default templates: {str(e)}")
            self.root.destroy()

    def setup_ui(self):
        """Setup the main user interface"""
        self.root.title("Automated ISD Document Generator")
        self.root.geometry("1920x1080")
        self.root.state("zoomed")

        # Main container
        main_frame = tb.Frame(self.root)
        main_frame.pack(fill=tk.BOTH, expand=True, padx=20, pady=20)

        # Left panel - Controls (store as self.control_frame)
        self.control_frame = tb.Frame(main_frame)
        self.control_frame.pack(side=tk.LEFT, fill=tk.Y, padx=10, pady=10)

        # Control buttons
        btn_data = tb.Button(self.control_frame, text="📂 Upload Data File",
                             command=self.upload_data_file)
        btn_data.pack(fill=tk.X, padx=10, pady=5)

        btn_output = tb.Button(self.control_frame, text="📁 Select Output Folder",
                               command=self.select_output_folder)
        btn_output.pack(fill=tk.X, padx=10, pady=5)

        btn_start = tb.Button(self.control_frame, text="🚀 Generate ISD Invoices",
                              bootstyle="success",
                              command=self.start_processing)
        btn_start.pack(fill=tk.X, padx=10, pady=20)

        # Add progress bar components (hidden initially)
        self.progress_frame = tb.Frame(self.control_frame)
        self.progress_label = tb.Label(self.progress_frame, text="Ready", bootstyle="info")
        self.progress_label.pack(fill=tk.X)

        self.progress_bar = tb.Progressbar(
            self.progress_frame,
            orient="horizontal",
            length=200,
            mode="determinate",
            bootstyle="success-striped"
        )
        self.progress_bar.pack(fill=tk.X, pady=5)
        self.progress_frame.pack_forget()  # Hide initially

        # Template status labels
        self.lbl_eligible_template = tb.Label(self.control_frame,
                                              text=f"✅ Eligible Template: {os.path.basename(self.eligible_template)}",
                                              bootstyle="success")
        self.lbl_eligible_template.pack(fill=tk.X, padx=10, pady=5)

        self.lbl_ineligible_template = tb.Label(self.control_frame,
                                                text=f"✅ Ineligible Template: {os.path.basename(self.ineligible_template)}",
                                                bootstyle="success")
        self.lbl_ineligible_template.pack(fill=tk.X, padx=10, pady=5)

        # Status labels
        self.lbl_data = tb.Label(self.control_frame, text="No Data File Loaded", bootstyle="secondary")
        self.lbl_data.pack(fill=tk.X, padx=10, pady=5)

        self.lbl_output = tb.Label(self.control_frame, text="No Output Folder Selected", bootstyle="secondary")
        self.lbl_output.pack(fill=tk.X, padx=10, pady=5)

        # Right panel - Data Preview
        preview_frame = tb.Frame(main_frame)
        preview_frame.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True, padx=10, pady=10)

        preview_label = tb.Label(preview_frame, text="Data Preview", bootstyle="primary")
        preview_label.pack(fill=tk.X, pady=5)

        # Create the treeview with proper scrollbars
        self.tree = self.create_treeview(preview_frame)

    def has_tax_amounts(self, row, is_eligible):
        """Check if row has any tax amounts for the given type (eligible/ineligible)"""
        prefix = "ELIGIBLE_" if is_eligible else "INELIGIBLE_"
        # Use the specific column names from your Excel structure
        tax_fields = ['CGST_AS_IGST', 'SGST_AS_IGST', 'CGST_AS_CGST', 'SGST_UTGST_AS_SGST_UTGST']

        for tax in tax_fields:
            col_name = prefix + tax
            if col_name in row and pd.notna(row[col_name]):
                try:
                    if float(row[col_name]) > 0:
                        return True
                except (ValueError, TypeError):
                    continue
        return False

    def start_processing(self):
        """Start the document generation process with organized output folders"""
        if not all([self.input_file, self.output_folder]):
            messagebox.showerror("Error", "Please select data file and output folder!")
            return

        try:
            # Verify input file and output folder
            if not os.path.exists(self.input_file):
                logging.error(f"Input file not found: {self.input_file}")
                messagebox.showerror("Error", "Input file not found!")
                return

            if not os.path.isdir(self.output_folder):
                logging.error(f"Output folder not found: {self.output_folder}")
                messagebox.showerror("Error", "Output folder not found!")
                return

            logging.info(f"Input: {self.input_file}")
            logging.info(f"Output: {self.output_folder}")

            # Show and initialize progress bar
            self.progress_frame.pack(fill=tk.X, padx=10, pady=(20, 5))
            self.progress_bar['value'] = 0
            self.progress_label.config(text="Preparing...")
            self.root.update_idletasks()

            data = read_excel_csv(self.input_file)
            if data is None:
                messagebox.showerror("Error", "Failed to read data file.")
                self.progress_frame.pack_forget()
                return

            # Create main output folders
            eligible_folder = os.path.join(self.output_folder, "Eligible")
            ineligible_folder = os.path.join(self.output_folder, "Ineligible")
            temp_docx_folder = os.path.join(self.output_folder, "TEMP_DOCX")

            try:
                os.makedirs(eligible_folder, exist_ok=True)
                os.makedirs(ineligible_folder, exist_ok=True)
                os.makedirs(temp_docx_folder, exist_ok=True)
            except PermissionError as pe:
                messagebox.showerror("Permission Error",
                                     f"Cannot create output folders:\n{str(pe)}\n"
                                     "Please choose a different output location.")
                return

            total_rows = len(data)
            success_count = 0

            for idx, row in data.iterrows():
                try:
                    # Update progress
                    progress = (idx + 1) / total_rows * 100
                    self.progress_bar['value'] = progress
                    self.progress_label.config(text=f"Processing row {idx + 1} of {total_rows}")
                    self.root.update_idletasks()

                    logging.info(f"\nProcessing row {idx}:")
                    logging.info(
                        f"Eligible amounts - CGST: {row['ELIGIBLE_CGST_AS_IGST']}, "
                        f"SGST: {row['ELIGIBLE_SGST_AS_IGST']}, "
                        f"IGST: {row['ELIGIBLE_IGST_AS_IGST']}"
                    )
                    logging.info(
                        f"Ineligible amounts - CGST: {row['INELIGIBLE_CGST_AS_IGST']}, "
                        f"SGST: {row['INELIGIBLE_SGST_AS_IGST']}, "
                        f"IGST: {row['INELIGIBLE_IGST_AS_IGST']}"
                    )

                    # Process both eligible and ineligible documents
                    for is_eligible in [True, False]:
                        if not self.has_tax_amounts(row, is_eligible):
                            logging.info(f"No {'eligible' if is_eligible else 'ineligible'} amounts found")
                            continue

                        # Set paths based on eligibility
                        if is_eligible:
                            output_pdf_folder = eligible_folder
                            prefix = "Eligible"
                            template_path = self.eligible_template
                        else:
                            output_pdf_folder = ineligible_folder
                            prefix = "Ineligible"
                            template_path = self.ineligible_template

                        # Generate document
                        try:
                            doc = Document(template_path)
                        except Exception as e:
                            logging.error(f"Failed to open template: {str(e)}")
                            continue

                        placeholders = scan_template_placeholders(template_path)
                        row_data = prepare_row_data(row, placeholders, is_eligible)

                        if not replace_all_placeholders(doc, row_data):
                            logging.error(f"Skipping row {idx} due to replacement errors")
                            continue

                        # Save temporary DOCX
                        invoice_num = str(row.get('INVOICE_NUMBER', idx + 1)).strip()
                        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                        docx_filename = f"{prefix}_ISD_{invoice_num}_{timestamp}.docx"
                        docx_path = os.path.join(temp_docx_folder, docx_filename)

                        try:
                            doc.save(docx_path)
                        except Exception as e:
                            logging.error(f"Failed to save DOCX: {str(e)}")
                            continue

                        # Convert to PDF in appropriate folder
                        pdf_filename = f"{prefix}_ISD_{invoice_num}_{timestamp}.pdf"
                        pdf_path = os.path.join(output_pdf_folder, pdf_filename)

                        try:
                            from docx2pdf import convert
                            convert(docx_path, pdf_path)
                            success_count += 1
                            logging.info(f"Generated {pdf_filename}")
                        except Exception as e:
                            logging.error(f"PDF conversion failed: {str(e)}")
                            continue

                        # Delete temporary DOCX
                        try:
                            os.remove(docx_path)
                        except Exception as e:
                            logging.error(f"Failed to delete temp DOCX: {str(e)}")

                except Exception as e:
                    logging.error(f"Error processing row {idx}: {str(e)}", exc_info=True)
                    continue

            # Clean up temporary folder
            try:
                if os.path.exists(temp_docx_folder):
                    if not os.listdir(temp_docx_folder):
                        os.rmdir(temp_docx_folder)
                    else:
                        logging.warning(f"Temporary folder not empty: {temp_docx_folder}")
            except Exception as e:
                logging.error(f"Error cleaning temp folder: {str(e)}")

            # Final progress update
            self.progress_bar['value'] = 100
            self.progress_label.config(text=f"Completed: {success_count} documents generated")
            self.root.update_idletasks()

            messagebox.showinfo("Success",
                                f"Processing complete!\n\n"
                                f"Eligible PDFs: {eligible_folder}\n"
                                f"Ineligible PDFs: {ineligible_folder}\n"
                                f"Total generated: {success_count}")

        except Exception as e:
            if hasattr(self, 'progress_label'):
                self.progress_label.config(text="Processing failed!", bootstyle="danger")
            messagebox.showerror("Error", f"Processing failed: {str(e)}")
            logging.error(f"Processing error: {str(e)}")

    def is_row_eligible(self, row):
        """Determine if row contains eligible or ineligible data"""
        eligible_cols = [
            'ELIGIBLE_IGST_AS_IGST', 'ELIGIBLE_CGST_AS_IGST',
            'ELIGIBLE_SGST_AS_IGST', 'ELIGIBLE_CGST_AS_CGST',
            'ELIGIBLE_SGST_UTGST_AS_SGST_UTGST'
        ]

        # Check if any eligible tax amount is > 0
        for col in eligible_cols:
            if col in row:
                try:
                    val = float(row[col]) if pd.notna(row[col]) else 0
                    if val > 0:
                        return True
                except (ValueError, TypeError):
                    continue
        return False

    def create_treeview(self, parent_frame):
        """Create and configure the Treeview widget with proper scrollbars"""
        # Container frame
        container = tb.Frame(parent_frame)
        container.pack(fill=tk.BOTH, expand=True)

        # Treeview widget
        tree = ttk.Treeview(container, selectmode="extended")

        # Vertical Scrollbar
        yscroll = ttk.Scrollbar(container, orient="vertical", command=tree.yview)
        yscroll.pack(side=tk.RIGHT, fill=tk.Y)
        tree.configure(yscrollcommand=yscroll.set)

        # Horizontal Scrollbar
        xscroll = ttk.Scrollbar(container, orient="horizontal", command=tree.xview)
        xscroll.pack(side=tk.BOTTOM, fill=tk.X)
        tree.configure(xscrollcommand=xscroll.set)

        # Pack treeview last
        tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        return tree

    def on_tree_right_click(self, event, tree):
        """Right-click menu to auto-resize columns"""
        menu = tk.Menu(self.root, tearoff=0)
        menu.add_command(label="Auto-size Columns", command=lambda: self.auto_size_columns(tree))
        menu.post(event.x_root, event.y_root)

    def auto_resize_columns(self):
        """Automatically resize columns to fit content"""
        for col in self.tree["columns"]:
            # Set minimum width based on header
            min_width = tk.font.Font().measure(col[:20]) + 20  # Add padding

            # Check all items for content width
            for item in self.tree.get_children():
                cell_value = str(self.tree.set(item, col))
                cell_width = tk.font.Font().measure(cell_value[:30]) + 20  # Limit check to 30 chars
                if cell_width > min_width:
                    min_width = cell_width

            # Set final column width
            self.tree.column(col, width=min_width)

    def display_data(self, data):
        """Display data in Treeview using first row for column width reference"""
        # Clear existing data
        self.tree.delete(*self.tree.get_children())

        # Set up columns
        self.tree["columns"] = list(data.columns)
        self.tree["show"] = "headings"

        # Add first row and use it for column width reference
        if len(data) > 0:
            first_row = data.iloc[0]

            # Configure columns based on first row values
            for col in data.columns:
                # Get header width
                header_width = tk.font.Font().measure(col) + 20  # Add padding

                # Get first row cell content width
                cell_value = str(first_row[col])
                cell_width = tk.font.Font().measure(cell_value) + 20  # Add padding

                # Use whichever is wider (header or first row content)
                col_width = max(header_width, cell_width)

                # Apply column configuration
                self.tree.heading(col, text=col)
                self.tree.column(col, width=col_width, stretch=False)  # Fixed width

            # Insert all rows (first row will match our column widths)
            for _, row in data.iterrows():
                self.tree.insert("", "end", values=list(row))
        else:
            # Empty dataset - just set up columns
            for col in data.columns:
                self.tree.heading(col, text=col)
                self.tree.column(col, width=tk.font.Font().measure(col) + 20, stretch=False)

        # Update the view
        self.tree.update_idletasks()

    def setup_menu(self):
        """Setup the menu bar"""
        menu_bar = tk.Menu(self.root)

        # File menu
        file_menu = tk.Menu(menu_bar, tearoff=0)
        file_menu.add_command(label="Upload Data File", command=self.upload_data_file)
        file_menu.add_command(label="Exit", command=self.root.quit)
        menu_bar.add_cascade(label="File", menu=file_menu)

        # Theme menu
        theme_menu = tk.Menu(menu_bar, tearoff=0)
        theme_options = {
            "darkly": "🌙 Dark",
            "journal": "📖 Light",
            "flatly": "📄 Flat",
            "cyborg": "🤖 Cyborg",
            "superhero": "🦸 Superhero",
            "minty": "🌿 Minty"
        }

        for theme, label in theme_options.items():
            theme_menu.add_command(label=label, command=lambda t=theme: self.change_theme(t))

        menu_bar.add_cascade(label="Theme", menu=theme_menu)

        self.root.config(menu=menu_bar)

    def change_theme(self, selected_theme):
        """Change the application theme"""
        self.root.style.theme_use(selected_theme)

    def upload_data_file(self):
        """Handle data file upload"""
        file_path = filedialog.askopenfilename(filetypes=[("Excel/CSV files", "*.xlsx;*.xls;*.csv")])
        if file_path:
            self.input_file = file_path
            self.lbl_data.config(text=f"📂 {os.path.basename(file_path)} Loaded")
            logging.info(f"Data file loaded: {file_path}")

            try:
                self.current_data = read_excel_csv(file_path)
                if self.current_data is not None:
                    self.display_data(self.current_data)
                    messagebox.showinfo("Success", "Data file loaded and displayed successfully!")
                else:
                    messagebox.showerror("Error", "Failed to read data file.")
            except Exception as e:
                messagebox.showerror("Error", f"Failed to load data: {str(e)}")
                logging.error(f"Data loading error: {str(e)}")

    def select_output_folder(self):
        """Handle output folder selection"""
        folder = filedialog.askdirectory()
        if folder:
            self.output_folder = folder
            self.lbl_output.config(text=f"📁 Output Folder: {folder}")
            logging.info(f"Output folder selected: {folder}")

    def setup_template_access(self):
        """Setup functionality for accessing the Excel template"""
        # Add button to UI
        self.template_button = tb.Button(
            self.control_frame,
            text="📊 Get Excel Template",
            command=self.provide_excel_template,
            bootstyle="info"
        )
        self.template_button.pack(fill=tk.X, padx=10, pady=5)

    def provide_excel_template(self):
        """Provide the Excel template to the user"""
        try:
            # Get the template from package resources
            source_path = self.get_template_path()

            # Determine where to save it
            save_path = filedialog.asksaveasfilename(
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")],
                title="Save Excel Template As",
                initialfile="ISD_Input_Template.xlsx"
            )

            if save_path:  # User didn't cancel
                shutil.copy(source_path, save_path)
                messagebox.showinfo(
                    "Template Saved",
                    f"Excel template saved to:\n{save_path}\n\n"
                    "Please use this format for your data."
                )
                # Optionally open the file
                if messagebox.askyesno("Open Template", "Open the template now?"):
                    os.startfile(save_path)

        except Exception as e:
            messagebox.showerror("Error", f"Could not provide template:\n{str(e)}")

    def get_template_path(self, template_type="excel"):
        """Get path to template whether running as script or executable"""
        try:
            # Validate template type
            template_files = {
                'excel': 'ISD_Input_Template.xlsx',
                'eligible': 'eligible_template.docx',
                'ineligible': 'ineligible_template.docx'
            }

            if template_type not in template_files:
                raise ValueError(f"Invalid template type. Must be one of: {list(template_files.keys())}")

            filename = template_files[template_type]

            # Try different locations
            possible_paths = []

            # 1. PyInstaller bundle location
            if getattr(sys, 'frozen', False):
                base_path = sys._MEIPASS
                possible_paths.append(os.path.join(base_path, "templates", filename))

            # 2. Development location (relative to script)
            possible_paths.append(os.path.join(os.path.dirname(__file__), "templates", filename))

            # 3. Current working directory
            possible_paths.append(os.path.join(os.getcwd(), "templates", filename))

            # 4. User's home directory
            possible_paths.append(os.path.join(os.path.expanduser("~"), "templates", filename))

            # Check each possible path
            for path in possible_paths:
                if os.path.exists(path):
                    return path

            raise FileNotFoundError(
                f"Could not locate {filename} in any of these locations:\n"
                + "\n".join(possible_paths)
            )

        except Exception as e:
            logging.error(f"Error finding template: {str(e)}")
            raise


# Initialize and run the application
if __name__ == "__main__":
    theme = "darkly" if darkdetect.isDark() else "journal"
    root = tb.Window(themename=theme)
    app = DocumentFillerApp(root)
    root.mainloop()