from docx import Document
import logging

# Set up logging
logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
logger = logging.getLogger(__name__)

def replace_placeholder_in_runs(paragraph, placeholder, replacement):
    """
    Replaces placeholders in a paragraph's runs while preserving formatting.

    Args:
        paragraph: A docx paragraph object.
        placeholder (str): The placeholder text to replace.
        replacement (str): The text to replace the placeholder with.
    """
    for run in paragraph.runs:
        if placeholder in run.text:
            run.text = run.text.replace(placeholder, replacement)

def map_data_to_docx(template_path, data, row_index=0):
    """
    Replaces placeholders in a DOCX template with data from a DataFrame.

    Args:
        template_path (str): Path to the DOCX template.
        data (pd.DataFrame): DataFrame containing the data.
        row_index (int): Index of the row to use for mapping (default: 0).

    Returns:
        Document: A docx.Document object with placeholders replaced, or None if an error occurs.
    """
    try:
        # Validate inputs
        if not template_path.endswith(".docx"):
            logger.error("Invalid template file format. Please provide a .docx file.")
            return None
        if data.empty:
            logger.error("The DataFrame is empty.")
            return None
        if row_index >= len(data):
            logger.error(f"Row index {row_index} is out of bounds.")
            return None

        # Load the DOCX template
        logger.info(f"Loading template: {template_path}")
        doc = Document(template_path)

        # Replace placeholders in each paragraph
        for para in doc.paragraphs:
            for col in data.columns:
                placeholder = f"{{{{{col.strip()}}}}}"
                if placeholder in para.text:
                    replace_placeholder_in_runs(para, placeholder, str(data[col].iloc[row_index]))

        logger.info("Successfully mapped data to template.")
        return doc

    except Exception as e:
        logger.error(f"Error mapping data to template: {e}")
        return None